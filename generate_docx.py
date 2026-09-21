from docx import Document
from docx.shared import Pt
import os

doc = Document()

# Title
doc.add_heading('Actualización para Issue #31', 0)

# Section 1
doc.add_heading('1. Cronograma con Fechas Reales (Para Sección 5.1)', level=1)
doc.add_paragraph('Reemplazar la tabla genérica de la sección 5.1 con esta (basada en el inicio de clases del 3 de agosto):')

# Table for Schedule
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Semana'
hdr_cells[1].text = 'Fechas Reales'
hdr_cells[2].text = 'Actividades Sugeridas (Fase del Proyecto)'

data = [
    ('Semana 1', '3 de agosto – 9 de agosto', 'Levantamiento de requerimientos y definición del proyecto.'),
    ('Semana 2', '10 de agosto – 16 de agosto', 'Diagnóstico de la empresa y estado del arte.'),
    ('Semana 3', '17 de agosto – 23 de agosto', 'Diseño de arquitectura y diagramas (MER, flujos).'),
    ('Semana 4', '24 de agosto – 30 de agosto', 'Configuración inicial del entorno y Docker (auth-proxy, inventory-service).'),
    ('Semana 5', '31 de agosto – 6 de septiembre', 'Implementación de Keycloak / Entra ID y pruebas de login.'),
    ('Semana 6', '7 de septiembre – 13 de septiembre', 'Integración del panel visual e inventario.'),
    ('Semana 7', '14 de septiembre – 20 de septiembre', 'Análisis de Calidad y Seguridad (SonarQube, OWASP ZAP).'),
    ('Semana 8', '21 de septiembre – 27 de septiembre', 'Ajustes finales, documentación y entrega del primer corte.')
]

for sem, fechas, act in data:
    row_cells = table.add_row().cells
    row_cells[0].text = sem
    row_cells[1].text = fechas
    row_cells[2].text = act

# Section 2
doc.add_paragraph()
doc.add_heading('2. Nuevos Riesgos (Para Sección 5.2 - Matriz de Riesgos)', level=1)
doc.add_paragraph('Agregar estos riesgos a los ya existentes en la sección 5.2:')

# Risks list
risks = [
    ('Riesgo 4: Limitaciones de hardware con Docker', 'Media', 'Alto', 'Incompatibilidad o falta de recursos (RAM/CPU) para ejecutar múltiples contenedores de Docker en los equipos de desarrollo.', 'Optimizar el docker-compose.yml para consumir menos recursos y asegurar requisitos mínimos de hardware.'),
    ('Riesgo 5: Conectividad con el tenant de Entra ID', 'Alta', 'Alto', 'Dificultades técnicas o de red al intentar sincronizar el tenant real de la empresa por políticas restrictivas de TI.', 'Realizar pruebas de concepto (PoC) tempranas y solicitar permisos de administrador con anticipación.'),
    ('Riesgo 6: Curva de aprendizaje en SonarQube y ZAP', 'Media', 'Medio', 'Retrasos en las tareas de evaluación de calidad debido a falta de experiencia previa con estas herramientas.', 'Reservar tiempo adicional en el sprint para investigación técnica y revisión de la documentación oficial.')
]

for title, prob, imp, desc, mit in risks:
    p = doc.add_paragraph()
    p.add_run(title).bold = True
    
    p_details = doc.add_paragraph()
    p_details.add_run('Probabilidad: ').bold = True
    p_details.add_run(prob + ' | ')
    p_details.add_run('Impacto: ').bold = True
    p_details.add_run(imp + '\n')
    
    p_details.add_run('Descripción: ').bold = True
    p_details.add_run(desc + '\n')
    
    p_details.add_run('Mitigación: ').bold = True
    p_details.add_run(mit)
    p_details.style = 'List Bullet'

doc_path = os.path.join(r'C:\Users\PC\Documents\Bot\proyecto-de-grado', 'Actualizacion_Issue31.docx')
doc.save(doc_path)
print(f"File created successfully at {doc_path}")
